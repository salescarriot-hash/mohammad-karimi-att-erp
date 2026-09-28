from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from difflib import SequenceMatcher
from pathlib import Path


def _clean(v):
    if v is None:
        return None
    s = str(v).replace("\u200c", " ").strip()
    return re.sub(r"\s+", " ", s) or None


def _norm(v):
    s = _clean(v) or ""
    return re.sub(r"[^a-z0-9]+", "", s.lower())


def _num(v):
    if v is None:
        return None
    s = str(v).strip().replace(",", "").replace(" ", "")
    s = re.sub(r"[^0-9.\-]", "", s)
    try:
        return float(s) if s else None
    except ValueError:
        return None


def _date(v):
    if not v:
        return None
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def extract_text(file_name: str, data: bytes) -> tuple[str, list[str]]:
    warnings: list[str] = []
    suffix = Path(file_name).suffix.lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            pages = []
            for i, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                pages.append(f"[PAGE {i}]\n{text}")
            joined = "\n\n".join(pages).strip()
            if joined:
                return joined, warnings
            warnings.append("PDF has no usable text; OCR adapter is required for scanned content.")
            return "", warnings
        except Exception as exc:
            warnings.append(f"PDF extraction failed: {exc}")
            return "", warnings
    if suffix in {".docx"}:
        try:
            from docx import Document
            doc = Document(io.BytesIO(data))
            chunks = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(chunks), warnings
        except Exception as exc:
            warnings.append(f"DOCX extraction failed: {exc}")
            return "", warnings
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
            chunks = []
            for ws in wb.worksheets:
                chunks.append(f"[SHEET {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    vals = [_clean(x) for x in row]
                    if any(vals):
                        chunks.append(" | ".join(v or "" for v in vals))
            return "\n".join(chunks), warnings
        except Exception as exc:
            warnings.append(f"XLSX extraction failed: {exc}")
            return "", warnings
    if suffix == ".csv":
        try:
            text = data.decode("utf-8-sig", errors="replace")
            rows = csv.reader(io.StringIO(text))
            return "\n".join(" | ".join(x.strip() for x in row) for row in rows), warnings
        except Exception as exc:
            warnings.append(f"CSV extraction failed: {exc}")
            return "", warnings
    if suffix in {".txt", ".log", ".md"}:
        return data.decode("utf-8-sig", errors="replace"), warnings
    warnings.append(f"Unsupported supplier quote format: {suffix or 'unknown'}")
    return "", warnings


def _find_field(text: str, labels: list[str], confidence=0.86):
    label = "(?:" + "|".join(re.escape(x) for x in labels) + ")"
    pat = re.compile(rf"(?im)^\s*{label}\s*[:#\-]?\s*(.+?)\s*$")
    m = pat.search(text)
    if m:
        value = _clean(m.group(1))
        if value and "|" in value:
            value = next((x.strip() for x in value.split("|") if x.strip()), value)
        return {"value": value, "confidence": confidence, "source": "text"}
    # also handle inline labels such as Quote No: ABC
    pat = re.compile(rf"(?i){label}\s*[:#\-]\s*([^|,;\n]+)")
    m = pat.search(text)
    if m:
        return {"value": _clean(m.group(1)), "confidence": confidence - 0.05, "source": "text-inline"}
    return {"value": None, "confidence": 0.0, "source": "not-found"}


def _extract_supplier_name(text: str):
    return _find_field(text, ["supplier", "vendor", "seller", "company", "supplier name", "vendor name"], 0.82)


def extract_supplier_quote(file_name: str, data: bytes, rfq_lines=None):
    text, warnings = extract_text(file_name, data)
    fields = []
    field_defs = {
        "supplier": ["supplier", "vendor", "seller", "company", "supplier name", "vendor name"],
        "quote_number": ["quote no", "quotation no", "quotation number", "quote number", "offer no", "offer number"],
        "quote_date": ["quote date", "quotation date", "date"],
        "valid_until": ["valid until", "validity", "valid through", "offer valid until"],
        "currency": ["currency"],
        "incoterm": ["incoterm", "delivery terms", "terms of delivery"],
        "delivery_time": ["delivery time", "lead time", "delivery"],
        "payment_terms": ["payment terms", "payment"],
        "origin": ["origin", "country of origin"],
        "warranty": ["warranty", "guarantee"],
    }
    for name, labels in field_defs.items():
        f = _find_field(text, labels)
        if f["value"] is not None:
            if name in {"quote_date", "valid_until"}:
                f["value"] = _date(f["value"]) or f["value"]
            fields.append({"field_name": name, **f})

    lines = []
    # Prefer pipe/table-like rows. Header matching is intentionally conservative.
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("[PAGE ") or line.startswith("[SHEET "):
            continue
        parts = [p.strip() for p in line.split("|")]
        if len(parts) < 3:
            continue
        joined = " ".join(parts).lower()
        if any(x in joined for x in ["part number", "unit price", "description", "qty", "quantity"]):
            continue
        # Typical: line | P/N | description | qty | unit | price | total
        if not re.match(r"^\s*\d{1,4}\s*$", parts[0]):
            continue
        candidate = {
            "source": f"text row {parts[0]}",
            "line_number": int(parts[0]),
            "supplier_part_number": _clean(parts[1] if len(parts) > 1 else None),
            "description": _clean(parts[2] if len(parts) > 2 else None),
            "quantity": _num(parts[3] if len(parts) > 3 else None),
            "unit": _clean(parts[4] if len(parts) > 4 else None),
            "unit_price": _num(parts[5] if len(parts) > 5 else None),
            "total_price": _num(parts[6] if len(parts) > 6 else None),
            "condition": _clean(parts[7] if len(parts) > 7 else None),
            "brand": _clean(parts[8] if len(parts) > 8 else None),
            "manufacturer": _clean(parts[9] if len(parts) > 9 else None),
            "lead_time": _clean(parts[10] if len(parts) > 10 else None),
            "availability": _clean(parts[11] if len(parts) > 11 else None),
            "technical_compliance": "NOT_REVIEWED",
            "deviation": _clean(parts[12] if len(parts) > 12 else None),
            "confidence": 0.72,
        }
        if candidate["supplier_part_number"] or candidate["description"]:
            lines.append(candidate)

    # Simple line pattern for plain text documents.
    if not lines:
        row_pat = re.compile(r"(?im)^\s*(\d{1,4})[.)\-]\s*(\S+)\s+(.+?)\s+(\d+(?:\.\d+)?)\s+(?:([A-Za-z]+)\s+)?([0-9][0-9,]*(?:\.\d+)?)\s*$")
        for m in row_pat.finditer(text):
            lines.append({
                "source": f"text row {m.group(1)}", "line_number": int(m.group(1)),
                "supplier_part_number": m.group(2), "description": _clean(m.group(3)),
                "quantity": _num(m.group(4)), "unit": _clean(m.group(5)),
                "unit_price": _num(m.group(6)), "total_price": None,
                "condition": None, "brand": None, "manufacturer": None,
                "lead_time": None, "availability": None,
                "technical_compliance": "NOT_REVIEWED", "deviation": None, "confidence": 0.62,
            })

    # Match extracted lines to customer RFQ lines, but never confirm the match.
    for item in lines:
        best = None
        best_score = 0.0
        for rl in rfq_lines or []:
            p1 = _norm(item.get("supplier_part_number")); p2 = _norm(getattr(rl, "customer_part_number", None))
            d1 = _norm(item.get("description")); d2 = _norm(getattr(rl, "description_normalized", None) or getattr(rl, "description_original", None))
            score = 0.0
            if p1 and p2 and p1 == p2:
                score = 1.0
            elif d1 and d2:
                score = SequenceMatcher(None, d1, d2).ratio()
            if score > best_score:
                best_score, best = score, rl
        item["rfq_line_candidate_id"] = getattr(best, "id", None) if best_score >= 0.45 else None
        item["rfq_line_match_confidence"] = round(best_score, 4)
        item["needs_user_confirmation"] = True

    supplier = next((x for x in fields if x["field_name"] == "supplier"), None)
    return {
        "document_type": "SUPPLIER_QUOTE",
        "fields": fields,
        "supplier_name_candidate": supplier["value"] if supplier else None,
        "lines": lines,
        "warnings": warnings,
        "requires_user_confirmation": True,
        "source_file": file_name,
    }
