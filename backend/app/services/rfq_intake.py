import csv, io, re
from pathlib import Path
from typing import Any
from app.services.translation import enrich_field, contains_persian, normalize_persian


def extract_text(filename: str, data: bytes) -> tuple[str, str, list[str]]:
    ext = Path(filename).suffix.lower()
    warnings: list[str] = []
    if ext in {".txt", ".csv"}:
        return data.decode("utf-8-sig", errors="replace"), "TEXT/CSV", warnings
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
            if not text.strip():
                warnings.append("PDF contains no extractable text; OCR is required for scanned/image-only pages.")
            return text, "PDF-TEXT", warnings
        except Exception as e:
            warnings.append(f"PDF text extraction failed: {e}")
            return "", "PDF", warnings
    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(io.BytesIO(data))
            chunks = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    vals = [c.text.strip() for c in row.cells if c.text.strip()]
                    if vals:
                        chunks.append(" | ".join(vals))
            return "\n".join(chunks), "DOCX", warnings
        except Exception as e:
            warnings.append(f"DOCX extraction failed: {e}")
            return "", "DOCX", warnings
    if ext in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if vals:
                        rows.append(" | ".join(vals))
            return "\n".join(rows), "XLSX", warnings
        except Exception as e:
            warnings.append(f"XLSX extraction failed: {e}")
            return "", "XLSX", warnings
    warnings.append("Unsupported document format for baseline extraction; AI/OCR adapter is required.")
    return "", "UNSUPPORTED", warnings


def first_match(text: str, labels: list[str]) -> tuple[str | None, str | None]:
    for label in labels:
        m = re.search(rf"(?im)^\s*{re.escape(label)}\s*[:#\-]?\s*(.+?)\s*$", text)
        if m:
            return m.group(1).strip(), m.group(0).strip()
    return None, None


def _norm_header(value: str) -> str:
    value = normalize_persian(value).lower().strip()
    value = re.sub(r"[^a-z0-9\u0600-\u06FF]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def _header_kind(value: str) -> str | None:
    h = _norm_header(value)
    aliases = {
        "line_number": {"line", "line no", "line number", "item", "item no", "row", "ردیف", "شماره ردیف", "شماره"},
        "description_original": {"description", "item description", "material description", "part description", "details", "شرح", "شرح کالا", "شرح قطعه", "شرح فنی", "نام کالا", "نام قطعه", "کالا", "قطعه"},
        "manufacturer": {"manufacturer", "maker", "brand", "oem", "سازنده", "تولید کننده", "تولیدکننده", "برند"},
        "customer_part_number": {"part number", "part no", "p n", "pn", "customer part number", "p n o", "شماره فنی", "شماره قطعه", "کد فنی", "پارت نامبر", "p n قطعه"},
        "quantity": {"qty", "quantity", "requested qty", "required qty", "تعداد", "مقدار", "تعداد مورد نیاز"},
        "unit": {"unit", "uom", "unit of measure", "واحد", "واحد اندازه گیری", "واحد اندازهگیری"},
        "condition_required": {"condition", "required condition", "وضعیت", "شرایط", "وضعیت کالا", "وضعیت قطعه"},
        "technical_requirements": {"technical requirements", "technical specification", "specification", "specs", "مشخصات فنی", "مشخصات", "مشخصات فنی کالا", "مشخصات فنی قطعه"},
        "documents_required": {"documents required", "required documents", "مدارک مورد نیاز", "مدارک لازم", "مدارک"},
        "delivery_requirement": {"delivery requirement", "delivery", "شرایط تحویل", "تحویل", "زمان تحویل", "موعد تحویل"},
    }
    for kind, names in aliases.items():
        if h in names:
            return kind
    return None


def _parse_number(value: str | None) -> str | None:
    if value is None:
        return None
    value = normalize_persian(value).strip().replace(",", "")
    value = value.translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789"))
    return value if re.fullmatch(r"\d+(?:\.\d+)?", value) else None


def _tabular_lines(text: str, filename: str) -> list[dict[str, Any]]:
    rows = [r.strip() for r in text.splitlines() if r.strip()]
    if not rows:
        return []

    # Find a header row containing at least two recognised RFQ columns.
    header_idx = None
    mapping: dict[int, str] = {}
    for idx, raw in enumerate(rows[:30]):
        if "|" in raw:
            delimiter_pattern = r"\s*\|\s*"
        elif "\t" in raw:
            delimiter_pattern = r"\t"
        elif ";" in raw:
            delimiter_pattern = r"\s*;\s*"
        else:
            delimiter_pattern = r"\s*,\s*"
        cells = [c.strip() for c in re.split(delimiter_pattern, raw) if c.strip()]
        kinds = {i: _header_kind(c) for i, c in enumerate(cells)}
        if len([k for k in kinds.values() if k]) < 2 and "|" not in raw and "\t" not in raw and ";" not in raw and "," not in raw:
            alt_cells = [c.strip() for c in re.split(r"\s{2,}", raw) if c.strip()]
            alt_kinds = {i: _header_kind(c) for i, c in enumerate(alt_cells)}
            if len([k for k in alt_kinds.values() if k]) >= 2:
                cells, kinds = alt_cells, alt_kinds
        recognised = {i: k for i, k in kinds.items() if k}
        if len(recognised) >= 2 and any(k == "description_original" for k in recognised.values()):
            header_idx, mapping = idx, recognised
            break
    if header_idx is None:
        return []

    result = []
    next_line = 1
    for raw in rows[header_idx + 1:]:
        if "|" in raw:
            delimiter_pattern = r"\s*\|\s*"
        elif "\t" in raw:
            delimiter_pattern = r"\t"
        elif ";" in raw:
            delimiter_pattern = r"\s*;\s*"
        else:
            delimiter_pattern = r"\s*,\s*"
        cells = [c.strip() for c in re.split(delimiter_pattern, raw) if c.strip()]
        if len(cells) <= 1 and "|" not in raw and "\t" not in raw and ";" not in raw and "," not in raw:
            cells = [c.strip() for c in re.split(r"\s{2,}", raw) if c.strip()]
        if not cells:
            continue
        vals: dict[str, str] = {}
        for i, kind in mapping.items():
            if i < len(cells):
                vals[kind] = cells[i]
        desc = vals.get("description_original")
        if not desc:
            continue
        line_no = vals.get("line_number")
        try:
            line_number = int(float(line_no)) if line_no else next_line
        except ValueError:
            line_number = next_line
        next_line = max(next_line, line_number + 1)
        fields = []
        for field_name in ("description_original", "manufacturer", "customer_part_number", "quantity", "unit", "condition_required", "technical_requirements", "documents_required", "delivery_requirement"):
            value = vals.get(field_name)
            if value:
                confidence = 0.86 if field_name in {"description_original", "quantity", "customer_part_number"} else 0.80
                field = {"field_name": field_name, "value": value, "confidence": confidence,
                         "source": f"{filename}: {raw}", "needs_confirmation": True}
                fields.append(enrich_field(field))
        if fields:
            result.append({"line_number": line_number, "fields": fields})
    return result


def build_preview_from_text(filename: str, text: str, method: str = "RULE-BASED") -> dict[str, Any]:
    warnings: list[str] = []
    fields: list[dict[str, Any]] = []
    patterns = {
        "customer_rfq_number": ["RFQ No", "RFQ Number", "RFQ #", "Reference", "Reference No", "شماره استعلام", "شماره RFQ", "شماره درخواست", "شماره درخواست قیمت"],
        "request_date": ["Request Date", "RFQ Date", "Date", "تاریخ درخواست", "تاریخ استعلام", "تاریخ"],
        "required_date": ["Required Date", "Delivery Date", "Need Date", "تاریخ مورد نیاز", "تاریخ تحویل", "موعد تحویل"],
        "currency": ["Currency", "ارز", "واحد پول"],
        "delivery_location": ["Delivery Location", "Delivery Address", "Ship To", "محل تحویل", "آدرس تحویل", "محل تحویل کالا"],
        "incoterm": ["Incoterm", "Incoterms", "اینکوترمز", "شرایط تحویل بین المللی"],
        "title": ["Title", "Subject", "عنوان", "موضوع", "شرح درخواست"],
    }
    for name, labels in patterns.items():
        value, source = first_match(text, labels)
        if value:
            fields.append(enrich_field({"field_name": name, "value": value, "confidence": 0.72,
                           "source": f"{filename}: {source}", "needs_confirmation": True}))

    lines = _tabular_lines(text, filename)
    if not lines:
        for raw in text.splitlines():
            m = re.match(r"^\s*(\d+)\s*[|,;\t]\s*(.+)$", raw)
            if not m:
                continue
            n = int(m.group(1))
            rest = m.group(2).strip()
            parts = [x.strip() for x in re.split(r"[|,;\t]", rest) if x.strip()]
            line_fields = [enrich_field({"field_name": "description_original", "value": parts[0], "confidence": 0.65,
                            "source": f"{filename}: {raw.strip()}", "needs_confirmation": True})]
            if len(parts) > 1 and _parse_number(parts[-1]):
                line_fields.append(enrich_field({"field_name": "quantity", "value": parts[-1], "confidence": 0.70,
                                    "source": f"{filename}: {raw.strip()}", "needs_confirmation": True}))
            lines.append({"line_number": n, "fields": line_fields})

    # Translation is a distinct, reviewable stage between extraction and ERP facts.
    # Preserve the customer's original wording; use the English candidate only for
    # normalized matching/output after user confirmation.
    translation_required = 0
    for f in fields:
        if f.get("translation_status") == "TRANSLATION_REQUIRED":
            translation_required += 1
    for line in lines:
        for f in line.get("fields", []):
            if f.get("translation_status") == "TRANSLATION_REQUIRED":
                translation_required += 1
            if f.get("field_name") == "description_original" and f.get("translated_value"):
                f["normalized_value"] = f["translated_value"]
    if translation_required:
        warnings.append(f"{translation_required} extracted field(s) contain Persian text requiring AI translation and user confirmation.")
    if not fields and not lines:
        warnings.append("No reliable RFQ fields were identified. The document may use a non-standard layout; manual review is required.")
    return {"source_file_name": filename, "extraction_method": method, "fields": fields,
            "lines": lines, "warnings": warnings, "text": text,
            "translation_stage": "AFTER_EXTRACTION_BEFORE_RFQ_MAPPING",
            "translation_required_count": translation_required}



def build_preview(filename: str, data: bytes) -> dict[str, Any]:
    text, method, warnings = extract_text(filename, data)
    result = build_preview_from_text(filename, text, method)
    result["warnings"] = warnings + result["warnings"]
    return result
