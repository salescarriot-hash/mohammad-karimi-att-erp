from __future__ import annotations
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import re
from docx import Document

TOKEN_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
LINE_PREFIX = "line."


def _fmt(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return f"{value:,.2f}".rstrip("0").rstrip(".")
    return str(value)


def _replace_runs(paragraph, mapping):
    # Fast path: token contained in one run, preserving all other formatting.
    for run in paragraph.runs:
        for token, value in mapping.items():
            if token in run.text:
                run.text = run.text.replace(token, _fmt(value))
    # Fallback for tokens split across runs. This preserves the paragraph but may
    # normalize run formatting for that paragraph, which is preferable to leaving
    # an unresolved placeholder in an official template.
    text = paragraph.text
    if "{{" not in text:
        return
    replaced = text
    for token, value in mapping.items():
        replaced = replaced.replace(token, _fmt(value))
    if replaced != text:
        for run in paragraph.runs:
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = replaced
        else:
            paragraph.add_run(replaced)


def _replace_in_table(table, mapping):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                _replace_runs(p, mapping)
            for nested in cell.tables:
                _replace_in_table(nested, mapping)


def _all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                for nested in cell.tables:
                    yield from _table_paragraphs(nested)


def _table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _table_paragraphs(nested)


def _clone_row(table, row):
    new_tr = deepcopy(row._tr)
    row._tr.addnext(new_tr)
    return table.rows[table.rows.index(row) + 1]


def _row_text(row):
    return " ".join(cell.text for cell in row.cells)


def _line_mapping(line):
    return {
        "{{line.number}}": line.line_number,
        "{{line.description}}": line.description,
        "{{line.manufacturer}}": line.manufacturer,
        "{{line.part_number}}": line.part_number,
        "{{line.brand}}": line.brand,
        "{{line.quantity}}": line.quantity,
        "{{line.unit}}": line.unit,
        "{{line.condition}}": line.condition,
        "{{line.unit_price}}": line.final_selling_price,
        "{{line.total}}": line.selling_total_price,
        "{{line.currency}}": line.currency,
        "{{line.delivery_time}}": line.delivery_time,
        "{{line.origin}}": line.origin,
        "{{line.warranty}}": line.warranty,
    }


def _expand_line_rows(table, lines):
    marker_rows = [row for row in list(table.rows) if any("{{line." in cell.text for cell in row.cells)]
    for marker in marker_rows:
        # Replace marker with first line and clone for subsequent lines.
        if not lines:
            marker._tr.getparent().remove(marker._tr)
            continue
        template_tr = deepcopy(marker._tr)
        for idx, line in enumerate(lines):
            row = marker if idx == 0 else _insert_row_after(table, marker, template_tr)
            mapping = _line_mapping(line)
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_runs(p, mapping)
            if idx > 0:
                marker = row


def _insert_row_after(table, after_row, template_tr):
    new_tr = deepcopy(template_tr)
    after_row._tr.addnext(new_tr)
    # Re-resolve through table rows so cell proxies are valid.
    for row in table.rows:
        if row._tr is new_tr:
            return row
    raise RuntimeError("Unable to create template line row")


def quotation_mapping(q, customer, contact=None):
    return {
        "{{quotation.number}}": q.quotation_number,
        "{{quotation.date}}": q.quotation_date,
        "{{quotation.valid_until}}": q.valid_until,
        "{{quotation.currency}}": q.currency,
        "{{quotation.incoterm}}": q.incoterm,
        "{{quotation.delivery_location}}": q.delivery_location,
        "{{quotation.payment_terms}}": q.payment_terms,
        "{{quotation.delivery_time}}": q.delivery_time,
        "{{quotation.warranty}}": q.warranty,
        "{{quotation.origin}}": q.origin,
        "{{quotation.price_basis}}": q.price_basis,
        "{{quotation.subtotal}}": q.subtotal,
        "{{quotation.discount}}": q.discount,
        "{{quotation.additional_cost}}": q.additional_cost,
        "{{quotation.grand_total}}": q.grand_total,
        "{{quotation.customer_notes}}": q.customer_notes,
        "{{customer.name}}": getattr(customer, "name", "") if customer else "",
        "{{customer.name_en}}": getattr(customer, "name_en", "") if customer else "",
        "{{customer.country}}": getattr(customer, "country", "") if customer else "",
        "{{customer.city}}": getattr(customer, "city", "") if customer else "",
        "{{customer.address}}": getattr(customer, "address", "") if customer else "",
        "{{contact.name}}": getattr(contact, "name", "") if contact else "",
        "{{contact.position}}": getattr(contact, "position", "") if contact else "",
        "{{contact.email}}": getattr(contact, "email", "") if contact else "",
        "{{contact.phone}}": getattr(contact, "phone", "") if contact else "",
    }


def inspect_template(path: Path):
    doc = Document(path)
    tokens = []
    for p in _all_paragraphs(doc):
        tokens.extend(TOKEN_RE.findall(p.text))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                tokens.extend(TOKEN_RE.findall(cell.text))
    unique = []
    for token in tokens:
        if token not in unique:
            unique.append(token)
    return {"tokens": unique, "line_tokens": [x for x in unique if x.startswith(LINE_PREFIX)], "paragraph_count": len(list(_all_paragraphs(doc))), "table_count": len(doc.tables)}


def render_quotation(template_path: Path, output_path: Path, q, lines, customer, contact=None):
    doc = Document(template_path)
    mapping = quotation_mapping(q, customer, contact)
    for p in doc.paragraphs:
        _replace_runs(p, mapping)
    for table in doc.tables:
        _expand_line_rows(table, lines)
        _replace_in_table(table, mapping)
    unresolved = []
    for p in _all_paragraphs(doc):
        unresolved.extend(TOKEN_RE.findall(p.text))
    if unresolved:
        unique = []
        for token in unresolved:
            if token not in unique:
                unique.append(token)
        raise ValueError(f"Template contains unsupported or unresolved placeholders: {unique}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path
