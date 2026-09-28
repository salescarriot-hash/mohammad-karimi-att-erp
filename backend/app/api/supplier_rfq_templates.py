from __future__ import annotations
from datetime import datetime, timezone
from pathlib import Path
import json
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from sqlalchemy import select
from sqlalchemy.orm import Session
from app.db.session import get_db
from app.models.rfq import RFQ, RFQLine
from app.models.supplier import Supplier
from app.services.supplier_rfq_template_engine import inspect_supplier_rfq_template, render_supplier_rfq

router = APIRouter(prefix="/templates/supplier-rfq", tags=["Supplier RFQ Template Engine"])
ROOT = Path(__file__).resolve().parents[3]
TEMPLATE_DIR = ROOT / "storage" / "templates" / "supplier-rfq"
GENERATED = ROOT / "storage" / "generated" / "supplier-rfq"
TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
GENERATED.mkdir(parents=True, exist_ok=True)
META = TEMPLATE_DIR / "templates.json"


def _load():
    if not META.exists():
        return []
    try:
        return json.loads(META.read_text(encoding="utf-8"))
    except Exception:
        return []


def _save(rows):
    META.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")


@router.post("/upload", status_code=201)
async def upload(file: UploadFile = File(...)):
    name = Path(file.filename or "supplier_rfq.xlsx").name
    if not name.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Supplier RFQ template must be .xlsx or .xlsm")
    content = await file.read()
    if not content:
        raise HTTPException(400, "Empty template file")
    template_id = datetime.now(timezone.utc).strftime("srfq-%Y%m%d%H%M%S%f")
    path = TEMPLATE_DIR / f"{template_id}{Path(name).suffix.lower()}"
    path.write_bytes(content)
    try:
        inspection = inspect_supplier_rfq_template(path)
    except Exception as exc:
        path.unlink(missing_ok=True)
        raise HTTPException(400, f"Invalid Excel template: {exc}")
    rows = _load()
    record = {"id": template_id, "name": name, "path": str(path.relative_to(ROOT)), "uploaded_at": datetime.now(timezone.utc).isoformat(), "inspection": inspection, "active": not any(x.get("active") for x in rows)}
    rows.append(record)
    _save(rows)
    return record


@router.get("")
def list_templates():
    return _load()


@router.post("/{template_id}/activate")
def activate(template_id: str):
    rows = _load()
    found = False
    for row in rows:
        row["active"] = row["id"] == template_id
        found = found or row["active"]
    if not found:
        raise HTTPException(404, "Template not found")
    _save(rows)
    return next(x for x in rows if x["id"] == template_id)



@router.post("/ensure-standard")
def ensure_standard_template():
    """Create/activate ATT's internal standard supplier RFQ workbook when no template was supplied yet."""
    from openpyxl import Workbook
    rows = _load()
    active = next((x for x in rows if x.get("active")), None)
    if active:
        return active
    template_id = datetime.now(timezone.utc).strftime("srfq-standard-%Y%m%d%H%M%S%f")
    path = TEMPLATE_DIR / f"{template_id}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "RFQ"
    headers = [
        ("A1", "ARTIN TURBINE TECHNOLOGY"), ("A2", "SUPPLIER REQUEST FOR QUOTATION"),
        ("A4", "Supplier"), ("B4", "{{supplier.name_en}}"), ("D4", "RFQ No."), ("E4", "{{rfq.number}}"),
        ("A5", "RFQ Date"), ("B5", "{{rfq.date}}"), ("D5", "Required Date"), ("E5", "{{rfq.required_date}}"),
        ("A6", "Currency"), ("B6", "{{rfq.currency}}"), ("D6", "Delivery Location"), ("E6", "{{rfq.delivery_location}}"),
        ("A7", "Incoterm"), ("B7", "{{rfq.incoterm}}"),
    ]
    for cell, value in headers: ws[cell] = value
    cols = ["Line", "Description", "Manufacturer", "Part Number", "Qty", "Unit", "Condition", "Technical Requirements", "Documents Required", "Delivery Requirement"]
    for i, value in enumerate(cols, 1): ws.cell(9, i).value = value
    tokens = ["{{line.number}}","{{line.description}}","{{line.manufacturer}}","{{line.part_number}}","{{line.quantity}}","{{line.unit}}","{{line.condition}}","{{line.technical_requirements}}","{{line.documents_required}}","{{line.delivery_requirement}}"]
    for i, value in enumerate(tokens, 1): ws.cell(10, i).value = value
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    ws["A1"].font = Font(bold=True, size=16); ws["A2"].font = Font(bold=True, size=12)
    for c in ws[9]: c.font = Font(bold=True); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c in ws[10]: c.alignment = Alignment(vertical="top", wrap_text=True)
    widths=[8,42,22,24,12,12,18,42,35,30]
    for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=w
    wb.save(path)
    record={"id":template_id,"name":"ATT Standard Supplier RFQ.xlsx","path":str(path.relative_to(ROOT)),"uploaded_at":datetime.now(timezone.utc).isoformat(),"inspection":inspect_supplier_rfq_template(path),"active":True,"source":"SYSTEM_STANDARD"}
    rows.append(record); _save(rows); return record

@router.post("/{template_id}/render/{rfq_id}/{supplier_id}")
def render(template_id: str, rfq_id: int, supplier_id: int, db: Session = Depends(get_db)):
    row = next((x for x in _load() if x["id"] == template_id), None)
    if not row:
        raise HTTPException(404, "Template not found")
    rfq = db.get(RFQ, rfq_id)
    supplier = db.get(Supplier, supplier_id)
    if not rfq:
        raise HTTPException(404, "RFQ not found")
    if not supplier:
        raise HTTPException(404, "Supplier not found")
    template_path = ROOT / row["path"]
    if not template_path.exists():
        raise HTTPException(404, "Template file is missing from storage")
    lines = db.scalars(select(RFQLine).where(RFQLine.rfq_id == rfq.id).order_by(RFQLine.line_number)).all()
    out = GENERATED / f"RFQ-{rfq.id}-SUP-{supplier.id}-{template_id}{template_path.suffix}"
    render_supplier_rfq(template_path, out, rfq, supplier, lines)
    return FileResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=out.name)
